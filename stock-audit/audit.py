#!/usr/bin/env python3
"""
India stock sheet auditor — mechanical rule engine.

Implements the locked rules in team/checklists/india-stock-sheet.md against a
downloaded .xlsx export of the India stock sheet. Reports findings as stock # +
exact cell + what is wrong versus expected.

READ-ONLY. This script never writes to the sheet and has no Google credentials.

Columns are resolved by HEADER TEXT, not by letter, so the same code works on NEW
(which has Category columns E-G) and OLD (which does not, shifting everything after D).

Usage:
    python3 stock-audit/audit.py <file.xlsx> --label india-new
    python3 stock-audit/audit.py <file.xlsx> --label india-new --full
    python3 stock-audit/audit.py <file.xlsx> --label india-new --json
"""

import argparse, hashlib, json, os, re, sys
from datetime import datetime, date

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is not installed. Run: pip install openpyxl")

STATE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "state")
APPROVALS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "approvals.yaml")

# ---------------------------------------------------------------- column map

FIELDS = {
    "sr":        ["sr. no.", "sr no.", "sr no", "sr."],
    "date":      ["date"],
    "design":    ["design"],
    "design_no": ["design number", "design no."],
    "location":  ["location"],
    "gold":      ["gold details", "gold detail"],
    "size":      ["inch size", "inch/size", "inch / size", "size"],
    "gross":     ["gross weight"],
    "net":       ["net weight"],
    "tdw":       ["total diamond weight"],
    "breakup":   ["diamond weight breakup", "stone weight breakup"],
    "pcs":       ["dia pcs.", "dia pcs", "stone pcs.", "stone pcs"],
    "total_pcs": ["total dia pcs.", "total dia pcs", "total stone pcs"],
    "pointers":  ["pointers", "pointer"],
    "shape":     ["shape"],
    "sieve":     ["sieve / size", "sieve/size", "sieve"],
    "party":     ["party name"],
    "code":      ["product code"],
    "dia_usd":   ["diamond price ($)"],
    "gold_usd":  ["gold price ($)"],
    "labor_usd": ["labor ($)", "labour ($)"],
    "total_usd": ["total ($)"],
    "dia_inr":   ["diamond price (₹)"],
    "gold_inr":  ["gold price (₹)"],
    "labor_inr": ["labor (₹)", "labour (₹)"],
    "total_inr": ["total (₹)"],
    "comments":  ["comments"],
}

ALLOWED_LOCATIONS = {
    "usa", "usa sell", "india", "india sell", "hk", "hk sell",
    "dxb", "dxb sell", "memo", "repair", "melt",
}

# Price List rate cells, verified 2026-09-07
GOLD_RATE = {"14": ("$R$2", "$S$2"), "18": ("$T$2", "$U$2")}
LABOR_DIAMOND = ("$R$3", "$S$3")
LABOR_POLKI = ("$R$4", "$S$4")


def sr_keys(v):
    """A Multi master Sr may cover several stock numbers, e.g. 'S1708C/S1709'.
    Return every key the group should be reachable under."""
    raw = str(v).strip()
    keys = [raw]
    if "/" in raw:
        keys += [p.strip() for p in raw.split("/") if p.strip() and p.strip() != raw]
    seen, out = set(), []
    for k in keys:
        if k not in seen:
            seen.add(k)
            out.append(k)
    return out


def as_date(v):
    """STOCK stores real dates; Multi often stores 'DD/MM/YYYY' text. Compare fairly."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    return None


def norm(v):
    return re.sub(r"\s+", " ", str(v)).strip().lower() if v is not None else ""


def sr_text(v):
    """OLD Sr numbers come back as floats: show 4, not 4.0."""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def find_header_row(ws, limit=6):
    """The header row is the first row with several of our known field names."""
    known = {a for aliases in FIELDS.values() for a in aliases}
    best, best_hits = 1, 0
    for r in range(1, min(limit, ws.max_row) + 1):
        hits = sum(
            1 for c in range(1, ws.max_column + 1)
            if norm(ws.cell(r, c).value) in known
        )
        if hits > best_hits:
            best, best_hits = r, hits
    return best, best_hits


def map_columns(ws, header_row):
    """field name -> column index, resolved by header text."""
    seen = {}
    for c in range(1, ws.max_column + 1):
        h = norm(ws.cell(header_row, c).value)
        if h and h not in seen:
            seen[h] = c
    out = {}
    for field, aliases in FIELDS.items():
        for a in aliases:
            if a in seen:
                out[field] = seen[a]
                break
    return out


# ---------------------------------------------------------------- helpers

def num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    if isinstance(v, str):
        s = v.replace(",", "").strip()
        try:
            return float(s)
        except ValueError:
            return None
    return None


def blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def cell_ref(cols, field, row):
    if field not in cols:
        return "?"
    return f"{get_column_letter(cols[field])}{row}"


def karat_of(gold_details):
    m = re.search(r"\b(14|18)\s*K", str(gold_details or ""), re.I)
    return m.group(1) if m else None


def is_earring(design, category):
    t = f"{design} {category}".lower()
    return any(w in t for w in ("earring", "ear ring", "stud", "e-ring"))


def looks_polki(*texts):
    return any("polki" in str(t or "").lower() for t in texts)


def has_vlookup(formula):
    return "vlookup" in str(formula or "").lower()


def is_formula(v):
    return isinstance(v, str) and v.startswith("=")


VLOOKUP_RANGE = re.compile(r"\$([A-Z]{1,2})\$(\d+)\s*:\s*\$([A-Z]{1,2})\$(\d+)")


# ---------------------------------------------------------------- audit

class Audit:
    def __init__(self, path, label):
        self.label = label
        self.wbf = openpyxl.load_workbook(path, data_only=False)  # formulas
        self.wbv = openpyxl.load_workbook(path, data_only=True)   # cached values
        self.findings = []
        self.questions = []
        self.stock_name = self._pick("stock")
        self.multi_name = self._pick("multiple dia", "multi")
        self.price_name = self._pick("price list")

    def _pick(self, *needles):
        for n in self.wbf.sheetnames:
            for needle in needles:
                if needle in n.lower():
                    return n
        return None

    def flag(self, sr, cell, problem, severity="ERROR"):
        self.findings.append(
            {"stock": sr or "?", "cell": cell, "problem": problem, "severity": severity}
        )

    def ask(self, sr, cell, question):
        self.questions.append({"stock": sr or "?", "cell": cell, "question": question})

    # ---- price list ----------------------------------------------------

    def load_price_list(self):
        """Product codes and the last populated row of each lookup column."""
        self.codes, self.list_end = set(), {}
        if not self.price_name:
            return
        ws = self.wbv[self.price_name]
        for col in ("B", "J"):
            last, gap = 0, 0
            for r in range(4, min(ws.max_row, 400) + 1):
                v = ws[f"{col}{r}"].value
                if blank(v):
                    gap += 1
                    # A run of blanks ends the lookup block. Anything after it is a
                    # separate block (rate labels and the like), not product codes.
                    if gap >= 3 and last:
                        break
                    continue
                gap = 0
                self.codes.add(norm(v))
                last = r
            self.list_end[col] = last

    def check_vlookup_caps(self):
        """A lookup range that ends before the price list does silently misses codes."""
        for col, letter in (("B", "B"), ("J", "J")):
            end = self.list_end.get(col, 0)
            if not end:
                continue
            caps = set()
            for sheet in (self.stock_name, self.multi_name):
                if not sheet:
                    continue
                ws = self.wbf[sheet]
                for row in ws.iter_rows():
                    for c in row:
                        if not is_formula(c.value) or "vlookup" not in c.value.lower():
                            continue
                        for m in VLOOKUP_RANGE.finditer(c.value):
                            if m.group(1) == letter:
                                caps.add(int(m.group(4)))
            for cap in sorted(caps):
                if cap < end:
                    self.flag(
                        "ALL", f"Price List {letter}{cap + 1}:{letter}{end}",
                        f"VLOOKUP ranges cap at row {cap} but the {letter}-column price "
                        f"list runs to row {end}. Codes on rows {cap + 1}-{end} are "
                        f"invisible to every lookup.", "ERROR",
                    )
                elif cap == end:
                    self.ask(
                        "ALL", f"Price List {letter}{end}",
                        f"The {letter}-column price list is exactly full to row {end}, "
                        f"the same row every VLOOKUP caps at. The next code added lands "
                        f"outside the range. Extend the ranges before adding one?",
                    )

    # ---- stock ---------------------------------------------------------

    def rows_of(self, sheet_name):
        wsf, wsv = self.wbf[sheet_name], self.wbv[sheet_name]
        hr, hits = find_header_row(wsf)
        cols = map_columns(wsf, hr)
        if hits < 8:
            raise SystemExit(
                f"Could not identify the header row on '{sheet_name}' "
                f"(best guess row {hr}, only {hits} known headers matched). "
                f"The sheet layout may have changed — ask Deval before auditing."
            )
        return wsf, wsv, cols, hr

    def priced_srs(self):
        """Stock numbers whose STOCK row carries any price value or formula."""
        wsf, wsv, cols, hr = self.rows_of(self.stock_name)
        out = set()
        money = [f for f in ("dia_usd", "gold_usd", "labor_usd", "total_usd",
                             "dia_inr", "gold_inr", "labor_inr", "total_inr") if f in cols]
        for r in range(hr + 1, wsf.max_row + 1):
            sr = wsv.cell(r, cols["sr"]).value
            if blank(sr):
                continue
            if any(not blank(wsf.cell(r, cols[f]).value) for f in money):
                out.add(sr_text(sr))
        return out

    def srs_on_date(self, want):
        """Stock numbers whose DATE column equals `want`, as the export renders it."""
        wsf, wsv, cols, hr = self.rows_of(self.stock_name)
        out = {}
        for r in range(hr + 1, wsf.max_row + 1):
            sr = wsv.cell(r, cols["sr"]).value
            if blank(sr) or "date" not in cols:
                continue
            if as_date(wsv.cell(r, cols["date"]).value) == want:
                out[sr_text(sr)] = r
        return out

    def audit_stock(self, only_srs=None):
        wsf, wsv, cols, hr = self.rows_of(self.stock_name)
        self.stock_cols = cols
        need = ("sr", "date", "design", "tdw", "net", "gross")
        missing = [f for f in need if f not in cols]
        if missing:
            raise SystemExit(f"STOCK is missing expected columns: {missing}")

        def V(r, f):
            return wsv.cell(r, cols[f]).value if f in cols else None

        def F(r, f):
            return wsf.cell(r, cols[f]).value if f in cols else None

        def C(r, f):
            return cell_ref(cols, f, r)

        for r in range(hr + 1, wsf.max_row + 1):
            sr = V(r, "sr")
            if blank(sr):
                continue
            sr = sr_text(sr)
            if only_srs is not None and sr not in only_srs:
                continue

            tdw = num(V(r, "tdw")) or 0.0
            gross, net = num(V(r, "gross")), num(V(r, "net"))
            sieve = norm(V(r, "sieve"))
            code = V(r, "code")
            is_mix = sieve == "mix" or norm(code) == "mix"

            # A row with a Sr number and nothing else is one problem, not seven.
            core = ("date", "design", "location", "gold", "gross", "net", "party")
            empty = [f for f in core if f in cols and blank(V(r, f))]
            if len(empty) >= 6:
                self.flag(sr, cell_ref(cols, "sr", r),
                          f"Row has a Sr number but is otherwise empty "
                          f"({', '.join(empty)} all blank)")
                continue

            # --- identity and basics
            if blank(V(r, "date")):
                self.flag(sr, C(r, "date"), "DATE blank — stock-in date is compulsory")
            if blank(V(r, "design")):
                self.flag(sr, C(r, "design"), "DESIGN blank")
            if tdw > 0 and blank(V(r, "design_no")):
                self.flag(sr, C(r, "design_no"),
                          f"Design Number blank while TDW = {tdw} (>0, so it is compulsory)")

            loc = V(r, "location")
            if blank(loc):
                self.flag(sr, C(r, "location"), "LOCATION blank")
            elif norm(loc) not in ALLOWED_LOCATIONS:
                self.ask(sr, C(r, "location"),
                         f"Location {str(loc).strip()!r} is not in the allowed list. "
                         f"New location, or a typo?")

            gold = V(r, "gold")
            if blank(gold):
                self.flag(sr, C(r, "gold"), "Gold Details blank")
            kt = karat_of(gold)
            if gold and not kt:
                self.flag(sr, C(r, "gold"),
                          f"Gold Details {str(gold).strip()!r} — no 14KT/18KT found")

            size = V(r, "size")
            if blank(size):
                if not is_earring(V(r, "design"), V(r, "category") if "category" in cols else ""):
                    self.flag(sr, C(r, "size"),
                              "INCH SIZE blank (only earrings/studs may be blank or NA)")
            else:
                s = str(size).strip()
                if re.match(r"^US\s+\d", s, re.I):
                    self.flag(sr, C(r, "size"),
                              f"Ring size {s!r} has a space — expected {s.replace(' ', '')!r}")

            if blank(V(r, "gross")):
                self.flag(sr, C(r, "gross"), "GROSS WEIGHT blank")
            if blank(V(r, "net")):
                self.flag(sr, C(r, "net"), "NET WEIGHT blank")

            if gross == 0:
                self.flag(sr, C(r, "gross"), "GROSS WEIGHT is zero")
            if gross is not None and net is not None and gross > 0:
                if tdw > 0 and net >= gross:
                    rel = "equals" if net == gross else "exceeds"
                    self.flag(sr, C(r, "net"),
                              f"Net {net} {rel} gross {gross} while TDW = {tdw}. "
                              f"With diamonds set, net must be less than gross")
                if tdw == 0 and net != gross:
                    self.flag(sr, C(r, "net"),
                              f"TDW = 0 but net {net} != gross {gross}. "
                              f"With no diamonds they should be equal")

            if blank(V(r, "party")):
                self.flag(sr, C(r, "party"), "PARTY NAME blank — manufacturer is compulsory")

            # --- diamonds
            if tdw > 0:
                if blank(code):
                    self.flag(sr, C(r, "code"), f"Product Code blank while TDW = {tdw}")
                elif not is_mix and self.codes and norm(code) not in self.codes:
                    self.flag(sr, C(r, "code"),
                              f"Product Code {str(code).strip()!r} is not on the Price List")

                pcs = num(V(r, "pcs"))
                if pcs in (None, 0):
                    self.flag(sr, C(r, "pcs"),
                              f"Diamond pcs blank or zero while TDW = {tdw}")

                if not is_mix:
                    bk = num(V(r, "breakup"))
                    if bk is not None and abs(bk - tdw) > 0.001:
                        self.flag(sr, C(r, "breakup"),
                                  f"Single row: breakup {bk} != TDW {tdw}")

                ptr_f, ptr_v = F(r, "pointers"), num(V(r, "pointers"))
                if "pointers" in cols and not is_mix:
                    if blank(ptr_f):
                        self.flag(sr, C(r, "pointers"), "POINTERS blank")
                    elif not is_formula(ptr_f):
                        self.flag(sr, C(r, "pointers"),
                                  f"POINTERS is hardcoded {ptr_f!r} — must be a formula "
                                  f"(TDW / pcs)")
                    elif pcs:
                        want = tdw / pcs
                        if ptr_v is not None and abs(ptr_v - want) > 0.005:
                            self.flag(sr, C(r, "pointers"),
                                      f"POINTERS shows {ptr_v} but TDW/pcs = "
                                      f"{tdw}/{pcs} = {want:.4f}")

                if blank(sieve):
                    self.flag(sr, C(r, "sieve"), "Sieve / Size blank while TDW > 0")

            # --- pricing (MIX rows are priced on the Multi tab, so skip them here)
            if not is_mix:
                self.check_pricing(sr, r, cols, wsf, kt, V(r, "design"), code)

        return cols

    def check_pricing(self, sr, r, cols, wsf, kt, design, code):
        def F(f):
            return wsf.cell(r, cols[f]).value if f in cols else None

        def C(f):
            return cell_ref(cols, f, r)

        priced = any(not blank(F(f)) for f in
                     ("dia_usd", "gold_usd", "labor_usd", "total_usd"))
        if not priced:
            return

        # P1 — diamond price must be a VLOOKUP, never a hardcoded multiplier
        for f, cur in (("dia_usd", "$"), ("dia_inr", "₹")):
            v = F(f)
            if blank(v):
                continue
            if not is_formula(v):
                self.flag(sr, C(f), f"Diamond price ({cur}) is a typed value {v!r}, not a formula")
            elif not has_vlookup(v):
                self.flag(sr, C(f),
                          f"Diamond price ({cur}) is hardcoded {v!r} — must VLOOKUP the "
                          f"Price List via the product code")

        # P2 — gold rate must match the karat
        if kt:
            want = GOLD_RATE[kt]
            for f, idx, cur in (("gold_usd", 0, "$"), ("gold_inr", 1, "₹")):
                v = F(f)
                if blank(v) or not is_formula(v):
                    continue
                if want[idx] not in v:
                    other = GOLD_RATE["18" if kt == "14" else "14"][idx]
                    wrong = f" — it references {other}" if other in v else ""
                    self.flag(sr, C(f),
                              f"Gold price ({cur}) does not use the {kt}KT rate "
                              f"{want[idx]}{wrong}")

        # P3 — labour rate: polki has its own rate
        polki = looks_polki(design, code)
        for f, idx, cur in (("labor_usd", 0, "$"), ("labor_inr", 1, "₹")):
            v = F(f)
            if blank(v) or not is_formula(v):
                continue
            if polki and LABOR_DIAMOND[idx] in v:
                self.flag(sr, C(f),
                          f"Polki piece but labour ({cur}) still on the diamond-set rate "
                          f"{LABOR_DIAMOND[idx]} — should be {LABOR_POLKI[idx]}")
            elif not polki and LABOR_POLKI[idx] in v:
                self.ask(sr, C(f),
                         f"Labour ({cur}) uses the polki rate {LABOR_POLKI[idx]} but the "
                         f"design does not say polki. Is this a polki piece?")

        # P4 — total must be the sum of its three parts
        for tot, parts, cur in (("total_usd", ("dia_usd", "gold_usd", "labor_usd"), "$"),
                                ("total_inr", ("dia_inr", "gold_inr", "labor_inr"), "₹")):
            v = F(tot)
            if blank(v):
                continue
            if not is_formula(v):
                self.flag(sr, C(tot), f"Total ({cur}) is a typed value {v!r}, not a formula")
                continue
            refs = {cell_ref(cols, p, r) for p in parts if p in cols}
            missing = [x for x in refs if x not in v.replace("$", "")]
            if missing and "sum" not in v.lower():
                self.flag(sr, C(tot),
                          f"Total ({cur}) {v!r} does not add {', '.join(sorted(refs))}")

    # ---- multi ---------------------------------------------------------

    def audit_multi(self, only_srs=None):
        if not self.multi_name:
            return {}
        wsf, wsv, cols, hr = self.rows_of(self.multi_name)
        groups, cur = {}, None
        for r in range(hr + 1, wsf.max_row + 1):
            sr = wsv.cell(r, cols["sr"]).value if "sr" in cols else None
            if not blank(sr):
                cur = sr_text(sr)
                entry = {"master": r, "children": [], "raw_sr": cur}
                for k in sr_keys(cur):
                    groups[k] = entry
            elif cur:
                if any(not blank(wsv.cell(r, c).value) for c in cols.values()):
                    groups[cur]["children"].append(r)

        seen_masters = set()
        for sr, g in groups.items():
            if only_srs is not None and sr not in only_srs:
                continue
            if g["master"] in seen_masters:
                continue
            seen_masters.add(g["master"])
            sr = g["raw_sr"]
            rows = [g["master"]] + g["children"]
            for r in rows:
                code = wsv.cell(r, cols["code"]).value if "code" in cols else None
                bk = num(wsv.cell(r, cols["breakup"]).value) if "breakup" in cols else None
                has_dia = any(
                    not blank(wsf.cell(r, cols[f]).value)
                    for f in ("dia_usd", "dia_inr") if f in cols
                )
                if has_dia and blank(code):
                    self.flag(sr, cell_ref(cols, "code", r),
                              "Multi Product Code blank while a diamond price is present")
                for f, cur_sym in (("dia_usd", "$"), ("dia_inr", "₹")):
                    if f not in cols:
                        continue
                    v = wsf.cell(r, cols[f]).value
                    if blank(v):
                        continue
                    if not is_formula(v):
                        self.flag(sr, cell_ref(cols, f, r),
                                  f"Multi diamond price ({cur_sym}) is a typed value {v!r}")
                    elif not has_vlookup(v):
                        self.flag(sr, cell_ref(cols, f, r),
                                  f"Multi diamond price ({cur_sym}) hardcoded {v!r} — "
                                  f"must VLOOKUP the Price List via the product code")
        self.multi_cols = cols
        self.multi_groups = groups
        return groups

    def cross_check_mix(self, only_srs=None):
        """MIX rows on STOCK must have a Multi group whose breakup and pcs reconcile."""
        if not self.multi_name or not hasattr(self, "multi_groups"):
            return
        wsf, wsv, cols, hr = self.rows_of(self.stock_name)
        mcols, mws = self.multi_cols, self.wbv[self.multi_name]
        for r in range(hr + 1, wsf.max_row + 1):
            sr = wsv.cell(r, cols["sr"]).value
            if blank(sr):
                continue
            sr = sr_text(sr)
            if only_srs is not None and sr not in only_srs:
                continue
            sieve = norm(wsv.cell(r, cols["sieve"]).value) if "sieve" in cols else ""
            code = norm(wsv.cell(r, cols["code"]).value) if "code" in cols else ""
            if sieve != "mix" and code != "mix":
                continue

            g = self.multi_groups.get(sr)
            if not g:
                self.flag(sr, cell_ref(cols, "sieve", r),
                          "Marked MIX on STOCK but there is no Multi group for this Sr")
                continue

            tdw = num(wsv.cell(r, cols["tdw"]).value) or 0.0
            tot_pcs = num(wsv.cell(r, cols["total_pcs"]).value) if "total_pcs" in cols else None
            def group_sum(field):
                if field not in mcols:
                    return 0.0
                # The master row holds the first sieve line, not only the identity.
                return sum(num(mws.cell(rr, mcols[field]).value) or 0.0
                           for rr in [g["master"]] + g["children"])

            child_bk = group_sum("breakup")
            child_pcs = group_sum("pcs")

            if tdw and abs(child_bk - tdw) > 0.005:
                self.flag(sr, cell_ref(cols, "tdw", r),
                          f"STOCK TDW {tdw} != sum of Multi breakup {child_bk:.3f} "
                          f"(difference {tdw - child_bk:+.3f})")
            if tot_pcs and abs(child_pcs - tot_pcs) > 0.5:
                self.flag(sr, cell_ref(cols, "total_pcs", r),
                          f"STOCK total pcs {tot_pcs:g} != sum of Multi pcs {child_pcs:g}")

            mdate = mws.cell(g["master"], mcols["date"]).value if "date" in mcols else None
            sdate = wsv.cell(r, cols["date"]).value
            md, sd = as_date(mdate), as_date(sdate)
            if md and sd and md != sd:
                self.flag(sr, cell_ref(mcols, "date", g["master"]),
                          f"Multi date {md.isoformat()} != STOCK date {sd.isoformat()}")

            sparty = norm(wsv.cell(r, cols["party"]).value) if "party" in cols else ""
            if sparty and "party" in mcols:
                parties = {norm(mws.cell(rr, mcols["party"]).value)
                           for rr in [g["master"]] + g["children"]}
                parties.discard("")
                if parties and sparty not in parties:
                    self.flag(sr, cell_ref(mcols, "party", g["master"]),
                              f"No Multi row names STOCK party {sparty!r} "
                              f"(Multi has {', '.join(sorted(parties))})")


# ---------------------------------------------------------------- state

def canon(v):
    """Normalise a cell so a fingerprint tracks CONTENT, not representation.

    Google's exporter and any round-trip through a spreadsheet tool can change how
    a value is written (float precision, date formatting, stray whitespace) without
    anyone editing anything. Without this, a harmless re-export makes every row look
    changed and the twice-daily batch floods.
    """
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return f"{v:.6f}".rstrip("0")
    if isinstance(v, int):
        return str(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def fingerprint(path, sheet):
    """Hash each STOCK row so a later run can tell what actually changed."""
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb[sheet]
    hr, _ = find_header_row(ws)
    cols = map_columns(ws, hr)
    out = {}
    for r in range(hr + 1, ws.max_row + 1):
        sr = ws.cell(r, cols["sr"]).value
        if blank(sr):
            continue
        vals = [canon(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        out[sr_text(sr)] = hashlib.md5("|".join(vals).encode()).hexdigest()[:16]
    return out


def load_approvals():
    """Signed-off exceptions. Returns (entries, warning).

    If the file cannot be read, return NO approvals and a warning. Failing open —
    suppressing nothing — is the safe direction: an over-reported finding wastes a
    minute, a silently suppressed one reaches a manufacturer's pay run.
    """
    if not os.path.exists(APPROVALS):
        return [], None
    try:
        import yaml
    except ImportError:
        return [], ("PyYAML is not installed, so approvals.yaml could not be read. "
                    "Nothing is being suppressed — approved rows will appear as "
                    "failures until `pip install pyyaml` is run.")
    try:
        with open(APPROVALS) as fh:
            data = yaml.safe_load(fh) or {}
        return data.get("approvals") or [], None
    except Exception as exc:
        return [], f"approvals.yaml could not be parsed ({exc}). Nothing is suppressed."


def apply_approvals(findings, approvals, priced_srs):
    """Split findings into (still active, suppressed by an approval)."""
    active, suppressed = [], []
    for f in findings:
        hit = None
        stocks = {str(f["stock"])} | {x.strip() for x in str(f["stock"]).split("/")}
        for a in approvals:
            if str(a.get("stock", "")).strip() not in stocks:
                continue
            if str(a.get("cell", "")).strip() != str(f["cell"]).strip():
                continue
            m = a.get("match")
            if m and str(m).lower() not in str(f["problem"]).lower():
                continue
            if a.get("void_if") == "priced" and str(a.get("stock", "")).strip() in priced_srs:
                # the row gained a price, so the reasoning behind the approval no longer holds
                f = dict(f, problem=f["problem"] +
                         "  [approval lapsed: this row is now priced]")
                hit = None
                break
            hit = a
            break
        (suppressed if hit else active).append(f if not hit else dict(f, approval=hit))
    return active, suppressed


def load_state(label):
    p = os.path.join(STATE_DIR, f"{label}.json")
    if os.path.exists(p):
        with open(p) as fh:
            return json.load(fh)
    return {}


def save_state(label, fps, standing=None):
    os.makedirs(STATE_DIR, exist_ok=True)
    p = os.path.join(STATE_DIR, f"{label}.json")
    with open(p, "w") as fh:
        json.dump({"checked_at": datetime.now().isoformat(timespec="seconds"),
                   "standing": standing or [],
                   "rows": fps}, fh, indent=1, sort_keys=True)
    return p


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--label", default="india-new")
    ap.add_argument("--full", action="store_true",
                    help="audit every row, not just rows new or changed since last run")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--no-save", action="store_true",
                    help="do not update the state file (dry run)")
    ap.add_argument("--date",
                    help="audit only rows whose DATE column is this YYYY-MM-DD, as the "
                         "export renders it")
    ap.add_argument("--show-approved", action="store_true",
                    help="also list findings suppressed by an approval")
    ap.add_argument("--passfail", action="store_true",
                    help="list every row in scope as PASS or FAIL, not just the failures")
    a = ap.parse_args()

    audit = Audit(a.xlsx, a.label)
    if not audit.stock_name:
        sys.exit("No STOCK tab found in this workbook.")
    audit.load_price_list()

    prev_state = load_state(a.label)
    prev = prev_state.get("rows", {})
    prev_standing = set(prev_state.get("standing", []))
    now = fingerprint(a.xlsx, audit.stock_name)
    changed = {sr for sr, h in now.items() if prev.get(sr) != h}
    new_srs = {sr for sr in now if sr not in prev}
    only = None if (a.full or not prev) else changed
    scope_rows = None
    if a.date:
        want = datetime.strptime(a.date, "%Y-%m-%d").date()
        scope_rows = audit.srs_on_date(want)
        only = set(scope_rows)
        if not only:
            print(f"No rows carry DATE {a.date}.")
            return

    audit.check_vlookup_caps()
    audit.audit_stock(only)
    audit.audit_multi(only)
    audit.cross_check_mix(only)

    approvals, appr_warning = load_approvals()
    suppressed = []
    if approvals:
        audit.findings, suppressed = apply_approvals(
            audit.findings, approvals, audit.priced_srs())
    if appr_warning:
        print(f"> WARNING: {appr_warning}\n")

    def qhash(q):
        return hashlib.md5(f"{q['stock']}|{q['cell']}|{q['question']}".encode()).hexdigest()[:12]

    all_standing = [qhash(q) for q in audit.questions]
    if not a.full:
        audit.questions = [q for q in audit.questions if qhash(q) not in prev_standing]

    if a.date:
        scope_desc = f"{len(only)} row(s) dated {a.date}"
    scope = (f"{len(only)} row(s) dated {a.date}" if a.date else
             "full sheet" if only is None else
             f"{len(changed)} row(s) new or changed since {load_state(a.label).get('checked_at', '?')}")
    result = {
        "label": a.label,
        "file": os.path.basename(a.xlsx),
        "run_at": datetime.now().isoformat(timespec="seconds"),
        "scope": scope,
        "rows_total": len(now),
        "rows_new": sorted(new_srs),
        "rows_changed": sorted(changed - new_srs),
        "findings": audit.findings,
        "questions": audit.questions,
        "suppressed_by_approval": suppressed,
    }

    if not a.no_save:
        # Do NOT let the watermark move past a row that is still failing. Otherwise a
        # reported-but-unfixed error is never mentioned again, and the batch quietly
        # stops nagging about exactly the rows that need chasing. Failing rows keep a
        # sentinel fingerprint so they resurface every run until someone edits them.
        failed = set()
        for f in audit.findings:
            st = str(f["stock"])
            failed.add(st)
            failed.update(x.strip() for x in st.split("/"))
        keep = {sr: ("OPEN-FINDING" if sr in failed else h) for sr, h in now.items()}
        result["held_open"] = sorted(sr for sr in now if sr in failed)
        result["state_file"] = save_state(a.label, keep, sorted(set(all_standing)))

    if a.json:
        print(json.dumps(result, indent=2, ensure_ascii=False))
        return

    if a.passfail:
        def owns(scope_sr, finding_stock):
            # a Multi finding may be reported under a compound Sr, "S1708C/S1709"
            return scope_sr == finding_stock or scope_sr in str(finding_stock).split("/")

        targets = scope_rows if scope_rows is not None else {s_: None for s_ in sorted(only or now)}
        by_sr = {}
        for s_ in targets:
            by_sr[s_] = [f for f in audit.findings if owns(s_, f["stock"])]
        npass = sum(1 for v in by_sr.values() if not v)
        print(f"# India stock audit — pass/fail — {a.label}")
        print(f"Scope: {scope}.")
        print(f"**{npass} PASS · {len(by_sr) - npass} FAIL** out of {len(by_sr)} rows.\n")
        print("| Stock # | Row | Result | Problems |")
        print("|---|---|---|---|")
        for s_ in sorted(by_sr, key=lambda x: (targets.get(x) or 0, x)):
            fs = by_sr[s_]
            row = targets.get(s_) or ""
            if not fs:
                print(f"| {s_} | {row} | PASS | — |")
            else:
                cells = "; ".join(f"{f['cell']}: {f['problem']}" for f in fs)
                print(f"| {s_} | {row} | **FAIL ({len(fs)})** | {cells} |")
        if a.show_approved and suppressed:
            print(f"\n## Suppressed by approval ({len(suppressed)})\n")
            print("| Stock # | Cell | Problem | Approved |")
            print("|---|---|---|---|")
            for f in suppressed:
                ap_ = f.get("approval", {})
                print(f"| {f['stock']} | {f['cell']} | {f['problem']} | "
                      f"{ap_.get('approved_by','?')} {ap_.get('approved_on','')} |")
        if audit.questions:
            print(f"\n## Questions for Deval\n")
            print("| Stock # | Cell | Question |")
            print("|---|---|---|")
            for q in audit.questions:
                print(f"| {q['stock']} | {q['cell']} | {q['question']} |")
        return

    print(f"# India stock audit — {a.label}")
    print(f"Scope: {scope}. {len(now)} stock rows in sheet.")
    if new_srs:
        print(f"New since last run: {', '.join(sorted(new_srs)[:20])}"
              + (" ..." if len(new_srs) > 20 else ""))
    print()
    if audit.findings:
        print(f"## Errors ({len(audit.findings)})\n")
        print("| Stock # | Cell | Problem |")
        print("|---|---|---|")
        for f in audit.findings:
            print(f"| {f['stock']} | {f['cell']} | {f['problem']} |")
    else:
        print("## Errors\n\nNone.")
    if audit.findings:
        from collections import Counter
        kinds = Counter(re.split(r"[.(\[]|\s+-\s+", f["problem"])[0].strip()[:60]
                        for f in audit.findings)
        print(f"\n## Summary by rule\n")
        print("| Count | Rule |")
        print("|---|---|")
        for k, n in kinds.most_common():
            print(f"| {n} | {k} |")
    if audit.questions:
        print(f"\n## Questions for Deval ({len(audit.questions)})\n")
        print("| Stock # | Cell | Question |")
        print("|---|---|---|")
        for q in audit.questions:
            print(f"| {q['stock']} | {q['cell']} | {q['question']} |")


if __name__ == "__main__":
    main()
